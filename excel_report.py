from __future__ import annotations

import json
from datetime import date, datetime, time
from decimal import Decimal
from io import BytesIO
import re
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

EXCEL_REPORT_SHEETS = (
    ("vehicles", "Vehicles", (
        "id", "plate_number", "make", "model", "year", "color", "status",
        "odometer", "registration_expiry", "insurance_expiry", "created_at",
    )),
    ("drivers", "Drivers", (
        "id", "name", "license_number", "phone", "email", "status", "created_at",
    )),
    ("assignments", "Assignments", (
        "id", "driver_id", "driver_name", "vehicle_id", "plate_number", "make",
        "model", "assigned_date", "returned_date", "notes",
    )),
    ("maintenance", "Maintenance", (
        "id", "vehicle_id", "plate_number", "service_type", "description",
        "cost", "performed_on",
    )),
    ("fuel_logs", "Fuel Logs", (
        "id", "vehicle_id", "plate_number", "fuel_type", "quantity",
        "price_per_liter", "total_cost", "logged_on",
    )),
    ("trips", "Trips", (
        "id", "record_date", "ism_no", "shipment_date", "time_in", "time_out",
        "origin", "destination", "vehicle_id", "plate_no", "plate_number",
        "make", "model", "driver_id", "driver_name", "load_details", "trip_fuel",
        "route", "start_odometer", "end_odometer", "trip_date", "notes",
    )),
)

EXCEL_TRIP_FIELDS = (
    "id", "ism_no", "shipment_date", "time_in", "time_out", "origin",
    "destination", "plate_number", "make", "model", "driver_name",
    "load_details", "trip_fuel", "route", "start_odometer", "end_odometer",
    "trip_date", "notes",
)

EXCEL_CURRENCY_FIELDS = {"cost", "price_per_liter", "total_cost"}
EXCEL_DATE_FIELDS = {
    "advance_date", "assigned_date", "created_at", "effective_date",
    "insurance_expiry", "logged_on", "payroll_date", "performed_on",
    "record_date", "registration_expiry", "returned_date", "shipment_date",
    "trip_date",
}


def _excel_cell_value(field: str, value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    if field in EXCEL_DATE_FIELDS and isinstance(value, str):
        try:
            return date.fromisoformat(value)
        except ValueError:
            return value
    if field in {"time_in", "time_out"} and isinstance(value, str):
        try:
            return time.fromisoformat(value)
        except ValueError:
            return value
    return value


def build_excel_report(report: dict[str, Any]) -> BytesIO:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Dashboard"
    workbook.properties.title = "Fleet Management System Report"
    workbook.properties.subject = "System-wide fleet operations report"
    title_fill = PatternFill("solid", fgColor="17365D")
    header_fill = PatternFill("solid", fgColor="1F4E78")
    alternate_fill = PatternFill("solid", fgColor="EAF1F8")

    def format_sheet(
        worksheet: Any,
        title: str,
        headers: list[str],
        rows: list[list[Any]],
    ) -> None:
        worksheet.sheet_view.showGridLines = False
        worksheet.merge_cells(
            start_row=1, start_column=1, end_row=1, end_column=max(1, len(headers))
        )
        title_cell = worksheet.cell(row=1, column=1, value=title)
        title_cell.font = Font(name="Aptos Display", size=16, bold=True, color="FFFFFF")
        title_cell.fill = title_fill
        title_cell.alignment = Alignment(vertical="center")
        worksheet.row_dimensions[1].height = 30
        worksheet.cell(
            row=2,
            column=1,
            value=f"Generated {datetime.now().astimezone().strftime('%Y-%m-%d %H:%M %Z')}",
        )
        worksheet.cell(row=2, column=1).font = Font(italic=True, color="64748B")

        header_row = 4
        for column_index, header in enumerate(headers, start=1):
            cell = worksheet.cell(
                row=header_row,
                column=column_index,
                value=header.replace("_", " ").title(),
            )
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        worksheet.row_dimensions[header_row].height = 30

        for row_index, values in enumerate(rows, start=header_row + 1):
            for column_index, value in enumerate(values, start=1):
                cell = worksheet.cell(row=row_index, column=column_index, value=value)
                if isinstance(value, str):
                    cell.data_type = "s"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if row_index % 2:
                    cell.fill = alternate_fill
                field = headers[column_index - 1]
                if field in EXCEL_CURRENCY_FIELDS:
                    cell.number_format = '"₱"#,##0.00;[Red]-"₱"#,##0.00'
                elif field == "quantity":
                    cell.number_format = "#,##0.00"
                elif field in EXCEL_DATE_FIELDS and isinstance(value, date):
                    cell.number_format = "yyyy-mm-dd"
                elif field in {"time_in", "time_out"} and isinstance(value, time):
                    cell.number_format = "hh:mm"

        if headers:
            last_column = len(headers)
            last_row = max(header_row, header_row + len(rows))
            last_column_letter = get_column_letter(last_column)
            worksheet.auto_filter.ref = f"A{header_row}:{last_column_letter}{last_row}"
            worksheet.freeze_panes = "A5"
            worksheet.print_title_rows = "1:4"
            worksheet.page_setup.orientation = "landscape"
            worksheet.page_setup.fitToWidth = 1
            worksheet.page_setup.fitToHeight = 0
            worksheet.sheet_properties.pageSetUpPr.fitToPage = True
            for column_index, header in enumerate(headers, start=1):
                values_for_width = [header.replace("_", " ").title()]
                values_for_width.extend(
                    str(row[column_index - 1]) if row[column_index - 1] is not None else ""
                    for row in rows
                )
                width = min(max(max(map(len, values_for_width), default=0) + 2, 12), 42)
                worksheet.column_dimensions[get_column_letter(column_index)].width = width

    summary_rows = [
        [key.replace("_", " ").title(), value]
        for key, value in report["dashboard"].items()
    ]
    format_sheet(summary, "Fleet Management System - Dashboard", ["metric", "value"], summary_rows)
    for row in range(5, summary.max_row + 1):
        metric = summary.cell(row=row, column=1).value
        if metric in {"Total Fuel Spend", "Total Maintenance"}:
            summary.cell(row=row, column=2).number_format = '"₱"#,##0.00;[Red]-"₱"#,##0.00'

    for data_key, sheet_title, base_fields in EXCEL_REPORT_SHEETS:
        records = report[data_key]
        headers = list(base_fields)
        for record in records:
            headers.extend(
                field for field in record
                if field not in headers and field not in {"driver", "vehicle"}
            )
        rows = [
            [_excel_cell_value(field, record.get(field)) for field in headers]
            for record in records
        ]
        format_sheet(workbook.create_sheet(sheet_title), sheet_title, headers, rows)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def build_trip_excel_report(
    trips: list[dict[str, Any]],
    group_by: str | None = None,
) -> BytesIO:
    workbook = Workbook()
    header_fill = PatternFill("solid", fgColor="1F4E78")
    alternate_fill = PatternFill("solid", fgColor="EAF1F8")
    group_labels = {
        "driver_name": "Driver",
        "plate_number": "Truck",
        "origin": "Origin",
        "destination": "Destination",
    }
    groups: dict[str, list[dict[str, Any]]] = {}
    if group_by:
        label = group_labels[group_by]
        for trip in trips:
            group_name = str(trip.get(group_by) or f"Unassigned {label.lower()}")
            groups.setdefault(group_name, []).append(trip)
        report_groups = sorted(
            groups.items(),
            key=lambda item: item[0].casefold(),
        )
    else:
        report_groups = [("Trip History", trips)]

    used_sheet_names: set[str] = set()
    for index, (group_name, group_trips) in enumerate(report_groups):
        if index == 0:
            worksheet = workbook.active
        else:
            worksheet = workbook.create_sheet()

        if group_by:
            base_name = group_name
        else:
            base_name = "Trip History"
        safe_name = "".join(
            "_" if character in "[]:*?/\\\\" else character
            for character in base_name
        ).strip("'")[:31] or "Trip History"
        sheet_name = safe_name
        suffix = 2
        while sheet_name.casefold() in used_sheet_names:
            suffix_text = f" ({suffix})"
            sheet_name = f"{safe_name[:31 - len(suffix_text)]}{suffix_text}"
            suffix += 1
        used_sheet_names.add(sheet_name.casefold())
        worksheet.title = sheet_name
        worksheet.sheet_view.showGridLines = False
        worksheet.merge_cells(
            start_row=1,
            start_column=1,
            end_row=1,
            end_column=len(EXCEL_TRIP_FIELDS),
        )
        report_title = (
            f"Fleet Management System - Trip History by {group_labels[group_by]}: {group_name}"
            if group_by
            else "Fleet Management System - Trip History"
        )
        title = worksheet.cell(row=1, column=1, value=report_title)
        title.font = Font(name="Aptos Display", size=16, bold=True, color="FFFFFF")
        title.fill = PatternFill("solid", fgColor="17365D")
        title.alignment = Alignment(vertical="center")
        worksheet.row_dimensions[1].height = 30
        worksheet.cell(
            row=2,
            column=1,
            value=f"Generated {datetime.now().astimezone().strftime('%Y-%m-%d %H:%M %Z')}",
        ).font = Font(italic=True, color="64748B")

        for column_index, field in enumerate(EXCEL_TRIP_FIELDS, start=1):
            cell = worksheet.cell(
                row=4,
                column=column_index,
                value=field.replace("_", " ").title(),
            )
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        worksheet.row_dimensions[4].height = 30

        for row_index, trip in enumerate(group_trips, start=5):
            for column_index, field in enumerate(EXCEL_TRIP_FIELDS, start=1):
                value = _excel_cell_value(field, trip.get(field))
                cell = worksheet.cell(row=row_index, column=column_index, value=value)
                if isinstance(value, str):
                    cell.data_type = "s"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if row_index % 2:
                    cell.fill = alternate_fill
                if field in EXCEL_DATE_FIELDS and isinstance(value, date):
                    cell.number_format = "yyyy-mm-dd"
                elif field in {"time_in", "time_out"} and isinstance(value, time):
                    cell.number_format = "hh:mm"

        last_column = get_column_letter(len(EXCEL_TRIP_FIELDS))
        last_row = max(4, 4 + len(group_trips))
        worksheet.auto_filter.ref = f"A4:{last_column}{last_row}"
        worksheet.freeze_panes = "A5"
        worksheet.print_title_rows = "1:4"
        worksheet.page_setup.orientation = "landscape"
        worksheet.page_setup.fitToWidth = 1
        worksheet.page_setup.fitToHeight = 0
        worksheet.sheet_properties.pageSetUpPr.fitToPage = True
        for column_index, field in enumerate(EXCEL_TRIP_FIELDS, start=1):
            values = [field.replace("_", " ").title()]
            values.extend(
                str(trip.get(field)) if trip.get(field) is not None else ""
                for trip in group_trips
            )
            worksheet.column_dimensions[get_column_letter(column_index)].width = min(
                max(max(map(len, values), default=0) + 2, 12),
                42,
            )

    if not report_groups:
        worksheet = workbook.active
        worksheet.title = "Trip History"
        worksheet.append(["No trips match the selected view"])

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def build_payroll_excel_report(
    report: dict[str, Any],
    trips: list[dict[str, Any]],
    drivers: list[dict[str, Any]],
    start_date: str,
    end_date: str,
    filters: dict[str, str],
) -> BytesIO:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Payroll Summary"
    workbook.properties.title = "Payroll Report"
    workbook.properties.subject = f"Payroll report for {start_date} to {end_date}"
    title_fill = PatternFill("solid", fgColor="334155")
    header_fill = PatternFill("solid", fgColor="475569")
    alternate_fill = PatternFill("solid", fgColor="F1F5F9")
    currency_format = '"₱"#,##0.00;[Red]-"₱"#,##0.00'
    period = f"Pay period: {start_date} to {end_date} | Status: {report.get('status', 'draft').title()}"
    filter_values = [
        f"{key.replace('_', ' ').title()}: {value}"
        for key, value in filters.items() if value
    ]
    filter_summary = "Active filters: " + ("; ".join(filter_values) if filter_values else "None")

    def write_title(worksheet: Any, title: str, columns: int) -> None:
        worksheet.sheet_view.showGridLines = False
        worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(1, columns))
        cell = worksheet.cell(row=1, column=1, value=title)
        cell.font = Font(name="Aptos Display", size=16, bold=True, color="FFFFFF")
        cell.fill = title_fill
        cell.alignment = Alignment(vertical="center")
        worksheet.row_dimensions[1].height = 30
        worksheet.cell(row=2, column=1, value=period).font = Font(color="475569", bold=True)
        worksheet.cell(row=3, column=1, value=filter_summary).font = Font(color="64748B", italic=True)

    def write_table(
        worksheet: Any,
        row_start: int,
        headers: list[str],
        rows: list[list[Any]],
        currency_columns: set[int] | None = None,
    ) -> None:
        currency_columns = currency_columns or set()
        for column, header in enumerate(headers, start=1):
            cell = worksheet.cell(row=row_start, column=column, value=header)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        worksheet.row_dimensions[row_start].height = 28
        for row_index, values in enumerate(rows, start=row_start + 1):
            for column, value in enumerate(values, start=1):
                cell = worksheet.cell(row=row_index, column=column, value=value)
                if isinstance(value, str):
                    cell.data_type = "s"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if row_index % 2:
                    cell.fill = alternate_fill
                if column in currency_columns and value is not None:
                    cell.number_format = currency_format
        if headers:
            last_row = max(row_start, row_start + len(rows))
            last_column = get_column_letter(len(headers))
            worksheet.auto_filter.ref = f"A{row_start}:{last_column}{last_row}"
            worksheet.freeze_panes = f"A{row_start + 1}"
            worksheet.print_title_rows = f"1:{row_start}"
            worksheet.page_setup.orientation = "landscape"
            worksheet.page_setup.fitToWidth = 1
            worksheet.page_setup.fitToHeight = 0
            worksheet.sheet_properties.pageSetUpPr.fitToPage = True
            for column, header in enumerate(headers, start=1):
                values_for_width = [header]
                values_for_width.extend(
                    str(row[column - 1]) if row[column - 1] is not None else ""
                    for row in rows
                )
                worksheet.column_dimensions[get_column_letter(column)].width = min(
                    max(max(map(len, values_for_width), default=0) + 2, 12), 40
                )

    def currency_sum(rows: list[dict[str, Any]], field: str) -> Decimal:
        return sum(
            (Decimal(str(row.get(field) or 0)).quantize(Decimal("0.01")) for row in rows),
            Decimal("0.00"),
        )

    summary_drivers = []
    for driver in drivers:
        driver_key = driver.get("driver_key") or f"id:{driver.get('driver_id')}"
        eligible_trip_count = sum(
            bool(trip.get("eligible"))
            and (trip.get("driver_key") or f"id:{trip.get('driver_id')}") == driver_key
            for trip in trips
        )
        summary_drivers.append({
            **driver,
            "visible_trip_count": eligible_trip_count,
        })

    eligible_trip_count = sum(bool(trip.get("eligible")) for trip in trips)
    totals = [
        ("Total Drivers", len(summary_drivers)),
        ("Total Trips", eligible_trip_count),
        ("Total Gross Payroll", currency_sum(summary_drivers, "gross_pay")),
        ("Total Cash Advances", currency_sum(summary_drivers, "cash_advances")),
        ("Total Overdraft Balance", currency_sum([
            {"value": driver.get("remaining_overdraft")
             if driver.get("remaining_overdraft") is not None
             else driver.get("overdraft_balance")}
            for driver in summary_drivers
        ], "value")),
        ("Total Net Payable", sum((
            Decimal(str(
                driver.get("final_net_pay")
                if driver.get("final_net_pay") is not None
                else max(Decimal(str(driver.get("net_before_overdraft") or 0)), Decimal("0"))
            )).quantize(Decimal("0.01"))
            for driver in summary_drivers
        ), Decimal("0.00"))),
    ]
    write_title(summary, "Payroll Report", 12)
    summary["A5"] = "REPORT TOTALS"
    summary["A5"].font = Font(bold=True, color="334155")
    for row_index, (label, value) in enumerate(totals, start=6):
        summary.cell(row=row_index, column=1, value=label)
        value_cell = summary.cell(row=row_index, column=2, value=value)
        if isinstance(value, Decimal):
            value_cell.number_format = currency_format
    statement_headers = [
        "Driver", "Eligible Trips", "Gross Earnings", "Cash Advances",
        "Claims", "Opening Overdraft", "New Overdraft", "Overdraft Deduction",
        "Outstanding Overdraft", "Net Before Overdraft", "Final Net Pay",
        "Remaining Amount Due",
    ]
    statement_rows = [[
        driver.get("driver_name"),
        driver["visible_trip_count"],
        Decimal(str(driver.get("gross_pay") or 0)),
        Decimal(str(driver.get("cash_advances") or 0)),
        "Coming Soon",
        Decimal(str(driver.get("opening_overdraft") or 0)),
        Decimal(str(driver.get("new_overdraft") or 0)),
        Decimal(str(driver.get("overdraft_deduction") or 0)),
        Decimal(str(
            driver.get("remaining_overdraft")
            if driver.get("remaining_overdraft") is not None
            else driver.get("overdraft_balance") or 0
        )),
        Decimal(str(driver.get("net_before_overdraft") or 0)),
        Decimal(str(driver.get("final_net_pay"))) if driver.get("final_net_pay") is not None else max(
            Decimal(str(driver.get("net_before_overdraft") or 0)), Decimal("0")
        ),
        Decimal(str(driver.get("remaining_amount_due") or 0)),
    ] for driver in summary_drivers]
    summary["A14"] = "DRIVER STATEMENTS"
    summary["A14"].font = Font(bold=True, color="334155")
    write_table(summary, 15, statement_headers, statement_rows, set(range(3, 13)) - {5})
    summary.column_dimensions["A"].width = 27
    summary.column_dimensions["B"].width = 16

    trip_headers = [
        "Date", "Driver Name", "ISM Number", "Origin", "Destination", "Trip Rate", "Review",
    ]
    trip_fields = [
        "payroll_date", "driver_name", "ism_no", "origin", "destination", "rate", "issue",
    ]
    trip_rows = [[
        Decimal(str(trip[field])) if field == "rate" and trip.get(field) is not None
        else trip.get(field)
        for field in trip_fields
    ] for trip in trips]
    trip_sheet = workbook.create_sheet("Trip Details")
    write_title(trip_sheet, "Payroll Trip Details", len(trip_headers))
    write_table(trip_sheet, 5, trip_headers, trip_rows, {6})

    related_driver_ids = {str(driver.get("driver_id")) for driver in summary_drivers}
    name_by_driver_id = {
        str(driver.get("driver_id")): driver.get("driver_name") for driver in summary_drivers
    }
    advances = [
        advance for advance in report.get("cash_advances", [])
        if str(advance.get("driver_id")) in related_driver_ids
    ]
    advance_sheet = workbook.create_sheet("Cash Advances")
    advance_headers = ["Driver", "Date", "Amount", "Reference", "Remarks"]
    advance_rows = [[
        name_by_driver_id.get(str(advance.get("driver_id")), advance.get("driver_id")),
        advance.get("advance_date"),
        Decimal(str(advance.get("amount") or 0)),
        advance.get("reference"),
        advance.get("remarks"),
    ] for advance in advances]
    write_title(advance_sheet, "Payroll Cash Advances", len(advance_headers))
    write_table(advance_sheet, 5, advance_headers, advance_rows, {3})

    overdraft_rows = [
        [driver.get("driver_name"), transaction.get("effective_date"),
         transaction.get("transaction_type"), Decimal(str(transaction.get("amount") or 0)),
         transaction.get("remarks")]
        for driver in summary_drivers
        for transaction in driver.get("overdraft_transactions", [])
    ]
    overdraft_sheet = workbook.create_sheet("Overdraft Ledger")
    overdraft_headers = ["Driver", "Effective Date", "Transaction", "Amount", "Remarks"]
    write_title(overdraft_sheet, "Payroll Overdraft Ledger", len(overdraft_headers))
    write_table(overdraft_sheet, 5, overdraft_headers, overdraft_rows, {4})

    issue_sheet = workbook.create_sheet("Review Items")
    write_title(issue_sheet, "Trips Requiring Review", len(trip_headers))
    issue_rows = [[
        trip.get(field) for field in trip_fields
    ] for trip in report.get("issues", [])]
    write_table(issue_sheet, 5, trip_headers, issue_rows)

    audit_sheet = workbook.create_sheet("Audit History")
    audit_headers = ["Date", "Action", "User", "Details"]
    audit_rows = [[
        entry.get("created_at"),
        entry.get("action"),
        entry.get("actor_id"),
        json.dumps(entry.get("details") or {}, ensure_ascii=False),
    ] for entry in report.get("audit_history", [])]
    write_title(audit_sheet, "Payroll Audit History", len(audit_headers))
    write_table(audit_sheet, 5, audit_headers, audit_rows)

    existing_sheet_names = {name.casefold() for name in workbook.sheetnames}
    for index, driver in enumerate(summary_drivers, start=1):
        driver_name = str(driver.get("driver_name") or "Unassigned Driver")
        sheet_base = re.sub(r"[\\/*?:\[\]]", "", f"Payslip {driver_name}").strip()
        sheet_base = sheet_base or f"Payslip Driver {index}"
        sheet_base = sheet_base[:31].rstrip()
        sheet_name = sheet_base
        suffix_index = 1
        while sheet_name.casefold() in existing_sheet_names:
            suffix = f" {driver.get('driver_id') or index}-{suffix_index}"
            sheet_name = f"{sheet_base[:31 - len(suffix)].rstrip()}{suffix}"
            suffix_index += 1
        existing_sheet_names.add(sheet_name.casefold())

        driver_sheet = workbook.create_sheet(sheet_name)
        driver_sheet.sheet_view.showGridLines = False
        driver_sheet.sheet_properties.pageSetUpPr.fitToPage = True
        driver_sheet.page_setup.orientation = "portrait"
        driver_sheet.page_setup.paperSize = driver_sheet.PAPERSIZE_LETTER
        driver_sheet.page_setup.fitToWidth = 1
        driver_sheet.page_setup.fitToHeight = 1
        driver_sheet.page_margins.left = 0.25
        driver_sheet.page_margins.right = 0.25
        driver_sheet.page_margins.top = 0.35
        driver_sheet.page_margins.bottom = 0.35
        driver_sheet.sheet_view.zoomScale = 85
        for column, width in {
            "A": 15, "B": 17, "C": 18, "D": 18, "E": 15, "F": 24,
        }.items():
            driver_sheet.column_dimensions[column].width = width

        def payslip_section(row: int, title: str, first_column: int, last_column: int) -> None:
            driver_sheet.merge_cells(
                start_row=row, start_column=first_column, end_row=row, end_column=last_column
            )
            cell = driver_sheet.cell(row=row, column=first_column, value=title)
            cell.font = Font(bold=True, color="FFFFFF", size=10)
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center")
            driver_sheet.row_dimensions[row].height = 20

        def payslip_label_value(
            row: int,
            label: str,
            value: Any,
            label_column: int,
            value_column: int,
            value_end_column: int | None = None,
            currency: bool = False,
        ) -> None:
            label_cell = driver_sheet.cell(row=row, column=label_column, value=label)
            label_cell.font = Font(bold=True, color="475569", size=9)
            label_cell.fill = alternate_fill
            label_cell.alignment = Alignment(vertical="center")
            if value_end_column is not None and value_end_column > value_column:
                driver_sheet.merge_cells(
                    start_row=row, start_column=value_column,
                    end_row=row, end_column=value_end_column,
                )
            value_cell = driver_sheet.cell(row=row, column=value_column, value=value)
            if isinstance(value, str):
                value_cell.data_type = "s"
            value_cell.font = Font(color="0F172A", size=9)
            value_cell.alignment = Alignment(vertical="center", horizontal="right" if currency else "left")
            if currency and value is not None:
                value_cell.number_format = currency_format
            driver_sheet.row_dimensions[row].height = 18

        title_cell = driver_sheet.cell(row=1, column=1, value="EMPLOYEE PAYSLIP")
        driver_sheet.merge_cells("A1:F1")
        title_cell.font = Font(name="Aptos Display", size=18, bold=True, color="FFFFFF")
        title_cell.fill = title_fill
        title_cell.alignment = Alignment(vertical="center")
        driver_sheet.row_dimensions[1].height = 34
        driver_sheet.merge_cells("A2:F2")
        company_cell = driver_sheet.cell(
            row=2, column=1, value="Good Luck Forwarding Systems Inc."
        )
        company_cell.font = Font(size=10, bold=True, color="334155")
        company_cell.alignment = Alignment(vertical="center")
        driver_sheet.merge_cells("A3:F3")
        period_cell = driver_sheet.cell(row=3, column=1, value=period)
        period_cell.font = Font(size=9, color="64748B")
        driver_sheet.merge_cells("A4:F4")
        filters_cell = driver_sheet.cell(row=4, column=1, value=filter_summary)
        filters_cell.font = Font(size=8, italic=True, color="64748B")

        payslip_section(6, "EMPLOYEE & PAY PERIOD DETAILS", 1, 6)
        payslip_label_value(7, "Employee", driver_name, 1, 2, 3)
        payslip_label_value(7, "Driver ID", driver.get("driver_id") or "—", 4, 5, 6)
        payslip_label_value(8, "Eligible trips", driver["visible_trip_count"], 1, 2, 3)
        payslip_label_value(8, "Payroll status", report.get("status", "draft").title(), 4, 5, 6)

        driver_key = driver.get("driver_key") or f"id:{driver.get('driver_id')}"
        driver_trips = [
            trip for trip in trips
            if (trip.get("driver_key") or f"id:{trip.get('driver_id')}") == driver_key
        ]
        payslip_section(10, "EARNINGS BREAKDOWN", 1, 6)
        earnings_headers = [
            "Date", "ISM Number", "Origin", "Destination", "Trip Rate", "Review",
        ]
        for column, header in enumerate(earnings_headers, start=1):
            cell = driver_sheet.cell(row=11, column=column, value=header)
            cell.font = Font(bold=True, color="FFFFFF", size=8)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        driver_sheet.row_dimensions[11].height = 25
        for row_index, trip in enumerate(driver_trips, start=12):
            values = [
                _excel_cell_value("payroll_date", trip.get("payroll_date")),
                trip.get("ism_no"),
                trip.get("origin"),
                trip.get("destination"),
                Decimal(str(trip["rate"])) if trip.get("rate") is not None else None,
                trip.get("issue") or ("Eligible" if trip.get("eligible") else "Needs review"),
            ]
            for column, value in enumerate(values, start=1):
                cell = driver_sheet.cell(row=row_index, column=column, value=value)
                if isinstance(value, str):
                    cell.data_type = "s"
                cell.font = Font(size=8, color="0F172A")
                cell.alignment = Alignment(
                    vertical="center",
                    horizontal="right" if column == 5 else "left",
                    wrap_text=column == 6,
                )
                if row_index % 2 == 0:
                    cell.fill = alternate_fill
                if column == 5 and value is not None:
                    cell.number_format = currency_format
                if column == 1 and isinstance(value, date):
                    cell.number_format = "yyyy-mm-dd"
            driver_sheet.row_dimensions[row_index].height = 18
        if not driver_trips:
            driver_sheet.merge_cells("A12:F12")
            empty_cell = driver_sheet.cell(
                row=12, column=1, value="No trip line items match the selected filters."
            )
            empty_cell.font = Font(size=8, italic=True, color="64748B")
            empty_cell.alignment = Alignment(vertical="center")
            driver_sheet.row_dimensions[12].height = 20

        deductions = [
            ("Cash advances", Decimal(str(driver.get("cash_advances") or 0))),
            ("Claims", "Coming Soon"),
            ("Opening overdraft", Decimal(str(driver.get("opening_overdraft") or 0))),
            ("New overdraft", Decimal(str(driver.get("new_overdraft") or 0))),
            ("Overdraft deduction", Decimal(str(driver.get("overdraft_deduction") or 0))),
            (
                "Outstanding overdraft",
                Decimal(str(
                    driver.get("remaining_overdraft")
                    if driver.get("remaining_overdraft") is not None
                    else driver.get("overdraft_balance") or 0
                )),
            ),
        ]
        net_summary = [
            ("Gross earnings", Decimal(str(driver.get("gross_pay") or 0))),
            ("Net before overdraft", Decimal(str(driver.get("net_before_overdraft") or 0))),
            (
                "Final net pay",
                Decimal(str(driver.get("final_net_pay")))
                if driver.get("final_net_pay") is not None
                else max(Decimal(str(driver.get("net_before_overdraft") or 0)), Decimal("0")),
            ),
            ("Remaining amount due", Decimal(str(driver.get("remaining_amount_due") or 0))),
        ]
        breakdown_start = max(13, 12 + len(driver_trips)) + 1
        payslip_section(breakdown_start, "DEDUCTIONS & ADJUSTMENTS", 1, 3)
        payslip_section(breakdown_start, "NET PAY SUMMARY", 4, 6)
        for offset in range(max(len(deductions), len(net_summary))):
            row_index = breakdown_start + offset + 1
            if offset < len(deductions):
                label, value = deductions[offset]
                payslip_label_value(
                    row_index, label, value, 1, 3, currency=isinstance(value, Decimal)
                )
            if offset < len(net_summary):
                label, value = net_summary[offset]
                payslip_label_value(
                    row_index, label, value, 4, 6, currency=True
                )
        advances = [
            advance for advance in report.get("cash_advances", [])
            if str(advance.get("driver_id")) == str(driver.get("driver_id"))
        ]
        detail_start = breakdown_start + max(len(deductions), len(net_summary)) + 2
        payslip_section(detail_start, "CASH ADVANCE DETAILS", 1, 6)
        advance_headers = ["Date", "Amount", "Reference", "Remarks"]
        for column, header in enumerate(advance_headers, start=1):
            cell = driver_sheet.cell(row=detail_start + 1, column=column, value=header)
            cell.font = Font(bold=True, color="FFFFFF", size=8)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row_index, advance in enumerate(advances, start=detail_start + 2):
            values = [
                _excel_cell_value("advance_date", advance.get("advance_date")),
                Decimal(str(advance.get("amount") or 0)),
                advance.get("reference"),
                advance.get("remarks"),
            ]
            for column, value in enumerate(values, start=1):
                cell = driver_sheet.cell(row=row_index, column=column, value=value)
                if isinstance(value, str):
                    cell.data_type = "s"
                cell.font = Font(size=8, color="0F172A")
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                if row_index % 2 == 0:
                    cell.fill = alternate_fill
                if column == 2:
                    cell.number_format = currency_format
                if column == 1 and isinstance(value, date):
                    cell.number_format = "yyyy-mm-dd"
            driver_sheet.merge_cells(
                start_row=row_index, start_column=4, end_row=row_index, end_column=6
            )
            driver_sheet.row_dimensions[row_index].height = 18
        if not advances:
            driver_sheet.merge_cells(
                start_row=detail_start + 2, start_column=1,
                end_row=detail_start + 2, end_column=6,
            )
            empty_cell = driver_sheet.cell(
                row=detail_start + 2, column=1, value="No cash advances recorded."
            )
            empty_cell.font = Font(size=8, italic=True, color="64748B")

        overdraft_transactions = driver.get("overdraft_transactions", [])
        overdraft_start = detail_start + max(len(advances), 1) + 3
        payslip_section(overdraft_start, "OVERDRAFT TRANSACTION HISTORY", 1, 6)
        overdraft_headers = ["Effective Date", "Transaction", "Amount", "Remarks"]
        for column, header in enumerate(overdraft_headers, start=1):
            cell = driver_sheet.cell(row=overdraft_start + 1, column=column, value=header)
            cell.font = Font(bold=True, color="FFFFFF", size=8)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row_index, transaction in enumerate(overdraft_transactions, start=overdraft_start + 2):
            values = [
                _excel_cell_value("effective_date", transaction.get("effective_date")),
                transaction.get("transaction_type"),
                Decimal(str(transaction.get("amount") or 0)),
                transaction.get("remarks"),
            ]
            for column, value in enumerate(values, start=1):
                cell = driver_sheet.cell(row=row_index, column=column, value=value)
                if isinstance(value, str):
                    cell.data_type = "s"
                cell.font = Font(size=8, color="0F172A")
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                if row_index % 2 == 0:
                    cell.fill = alternate_fill
                if column == 3:
                    cell.number_format = currency_format
                if column == 1 and isinstance(value, date):
                    cell.number_format = "yyyy-mm-dd"
            driver_sheet.merge_cells(
                start_row=row_index, start_column=4, end_row=row_index, end_column=6
            )
            driver_sheet.row_dimensions[row_index].height = 18
        if not overdraft_transactions:
            driver_sheet.merge_cells(
                start_row=overdraft_start + 2, start_column=1,
                end_row=overdraft_start + 2, end_column=6,
            )
            empty_cell = driver_sheet.cell(
                row=overdraft_start + 2, column=1, value="No overdraft transactions recorded."
            )
            empty_cell.font = Font(size=8, italic=True, color="64748B")

        final_row = driver_sheet.max_row
        driver_sheet.print_area = f"A1:F{final_row}"
        driver_sheet.sheet_properties.pageSetUpPr.autoPageBreaks = False

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
