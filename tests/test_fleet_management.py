import json
import os
import tempfile
import time
import unittest
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import patch

import app as web_app
from PIL import Image
from fleet_management import FleetManagementError, FleetManager, _supabase_project_url
from openpyxl import load_workbook
from payroll import calculate_driver_payroll, period_date, trip_rate
from werkzeug.exceptions import MethodNotAllowed


class PayrollCalculationTests(unittest.TestCase):
    def test_payroll_get_route_is_not_claimed_by_generic_resource_routes(self):
        adapter = web_app.app.url_map.bind("127.0.0.1")

        self.assertEqual(adapter.match("/api/payroll", method="GET")[0], "payroll_report")
        with self.assertRaises(MethodNotAllowed):
            adapter.match("/api/payroll", method="POST")
        self.assertEqual(adapter.match("/api/trips", method="POST")[0], "create_record")

    def test_known_route_rates_are_bidirectional_and_case_insensitive(self):
        routes = [
            ("DAV2", "PLAS", "200.00"),
            ("PLAS", "DAV2", "200.00"),
            ("DAV1", "DAV2", "200.00"),
            ("DAV2", "DAV1", "200.00"),
            ("DAV1", "PLAS", "200.00"),
            ("PLAS", "DAV1", "200.00"),
            ("GENSAN", "DAV2", "900.00"),
            ("DAV2", "GENSAN", "900.00"),
            ("GENSAN", "PLAS", "900.00"),
            ("PLAS", "GENSAN", "900.00"),
            ("GENSAN", "DAV1", "900.00"),
            ("DAV1", "GENSAN", "900.00"),
            ("gensan", "plas", "900.00"),
        ]
        for origin, destination, expected in routes:
            with self.subTest(origin=origin, destination=destination):
                self.assertEqual(str(trip_rate(origin, destination)), expected)
        self.assertIsNone(trip_rate(None, "DAV2"))

    def test_payroll_net_and_explicit_overdraft_settlement(self):
        trips = [{"rate": "900.00"}, {"rate": "200.00"}]
        advances = [{"amount": "300.00"}]
        before_settlement = calculate_driver_payroll(trips, advances, "500.00")
        self.assertEqual(before_settlement["gross_pay"], Decimal("1100.00"))
        self.assertEqual(before_settlement["net_before_overdraft"], Decimal("800.00"))
        self.assertEqual(before_settlement["overdraft_deduction"], Decimal("0.00"))
        settled = calculate_driver_payroll(
            trips, advances, "500.00", settle_overdraft=True
        )
        self.assertEqual(settled["overdraft_deduction"], Decimal("500.00"))
        self.assertEqual(settled["final_net_pay"], Decimal("300.00"))
        negative = calculate_driver_payroll(
            [{"rate": "200.00"}], [{"amount": "300.00"}], "500.00"
        )
        self.assertEqual(negative["net_before_overdraft"], Decimal("-100.00"))
        self.assertEqual(negative["remaining_amount_due"], Decimal("100.00"))
        self.assertEqual(negative["overdraft_deduction"], Decimal("0.00"))

    def test_payroll_period_dates_are_validated(self):
        self.assertEqual(period_date("2026-10-01", "Start date").isoformat(), "2026-10-01")
        with self.assertRaisesRegex(ValueError, "valid date"):
            period_date("not-a-date", "Start date")

    def test_payroll_report_uses_routes_without_requiring_trip_times(self):
        client = FakeSupabaseClient()
        manager = FleetManager(client=client)
        vehicle = manager.add_vehicle("PAY-001", "Toyota", "Hiace", 2022)
        driver = manager.add_driver("Payroll Driver", "PAY-DL-1")
        manager.add_trip(
            vehicle["id"], driver["id"], "DAV2 to PLAS", 0, 0,
            ism_no="ISM-200", shipment_date="2026-10-02",
            time_in="08:00", time_out="10:00", origin="DAV2", destination="PLAS",
        )
        manager.add_trip(
            vehicle["id"], driver["id"], "DAV1 to PLAS", 0, 0,
            ism_no="ISM-201", shipment_date="2026-10-03",
            time_in="08:00", time_out="10:00", origin="DAV1", destination="PLAS",
        )
        manager.add_trip(
            vehicle["id"], driver["id"], "DAV1 to DAV2", 0, 0,
            ism_no="ISM-202", shipment_date="2026-10-04",
            time_in="08:00", origin="DAV1", destination="DAV2",
        )
        manager.add_trip(
            vehicle["id"], driver["id"], "DAV1 to UNKNOWN", 0, 0,
            ism_no="ISM-203", shipment_date="2026-10-05",
            time_in="08:00", time_out="10:00", origin="DAV1", destination="UNKNOWN",
        )

        report = manager.payroll_report("2026-10-01", "2026-10-15")

        self.assertEqual(report["drivers"][0]["gross_pay"], "600.00")
        self.assertEqual(len([trip for trip in report["trips"] if trip["eligible"]]), 3)
        self.assertEqual(
            {trip["issue"] for trip in report["issues"]},
            {"Unconfigured Route"},
        )

    def test_payroll_matches_imported_trip_driver_name_when_driver_id_is_missing(self):
        client = FakeSupabaseClient()
        manager = FleetManager(client=client)
        vehicle = manager.add_vehicle("PAY-IMPORT", "Toyota", "Hiace", 2022)
        driver = manager.add_driver("Imported Driver", "PAY-DL-IMPORT")
        manager.add_trip(
            vehicle["id"], driver["id"], "DAV1 to PLAS", 0, 0,
            ism_no="ISM-IMPORT", shipment_date="2026-10-03",
            origin="DAV1", destination="PLAS",
        )
        client.tables["trips"][0]["driver_id"] = None
        client.tables["trips"][0]["driver_name"] = "  imported   DRIVER "

        report = manager.payroll_report("2026-10-01", "2026-10-15")

        self.assertEqual(report["trips"][0]["driver_id"], driver["id"])
        self.assertTrue(report["trips"][0]["eligible"])
        self.assertEqual(report["trips"][0]["rate"], "200.00")
        self.assertEqual(report["drivers"][0]["gross_pay"], "200.00")

    def test_payroll_matches_unique_imported_last_name_to_registered_driver(self):
        client = FakeSupabaseClient()
        manager = FleetManager(client=client)
        driver = manager.add_driver("Rey Madamba", "PAY-DL-MADAMBA")
        client.tables["trips"] = [{
            "id": 21,
            "driver_id": None,
            "driver_name": "  MADAMBA ",
            "ism_no": "ISM-MADAMBA",
            "shipment_date": "2026-10-03",
            "origin": "DAV1",
            "destination": "PLAS",
            "time_in": None,
            "time_out": None,
        }]

        report = manager.payroll_report("2026-10-01", "2026-10-15")

        self.assertEqual(report["trips"][0]["driver_id"], driver["id"])
        self.assertEqual(report["trips"][0]["driver_name"], "Rey Madamba")
        self.assertEqual(report["trips"][0]["driver_key"], f"id:{driver['id']}")
        self.assertEqual(report["drivers"][0]["gross_pay"], "200.00")
        self.assertEqual(report["drivers"][0]["total_trips"], 1)

    def test_payroll_does_not_guess_when_imported_last_name_is_ambiguous(self):
        client = FakeSupabaseClient()
        manager = FleetManager(client=client)
        manager.add_driver("Rey Madamba", "PAY-DL-MADAMBA-1")
        manager.add_driver("Jose Madamba", "PAY-DL-MADAMBA-2")
        client.tables["trips"] = [{
            "id": 22,
            "driver_id": None,
            "driver_name": "Madamba",
            "ism_no": "ISM-MADAMBA-AMBIGUOUS",
            "shipment_date": "2026-10-03",
            "origin": "DAV1",
            "destination": "PLAS",
            "time_in": None,
            "time_out": None,
        }]

        report = manager.payroll_report("2026-10-01", "2026-10-15")

        self.assertIsNone(report["trips"][0]["driver_id"])
        self.assertEqual(report["trips"][0]["driver_name"], "Madamba")
        self.assertEqual(report["trips"][0]["rate"], "200.00")
        self.assertEqual(
            [driver["gross_pay"] for driver in report["drivers"]],
            ["0.00", "0.00", "200.00"],
        )

    def test_payroll_calculates_unassigned_import_by_route_without_ism_or_times(self):
        client = FakeSupabaseClient()
        manager = FleetManager(client=client)
        manager.add_driver("Registered Driver", "PAY-DL-REGISTERED")
        client.tables["trips"] = [{
            "id": 12,
            "driver_id": None,
            "driver_name": "Manifest Only Driver",
            "ism_no": None,
            "shipment_date": "2026-10-03",
            "origin": "PLAS",
            "destination": "DAV1",
            "time_in": None,
            "time_out": None,
        }]

        report = manager.payroll_report("2026-10-01", "2026-10-15")

        self.assertEqual(report["trips"][0]["rate"], "200.00")
        self.assertTrue(report["trips"][0]["eligible"])
        self.assertEqual(report["trips"][0]["issue"], None)
        self.assertEqual(report["drivers"][1]["driver_name"], "Manifest Only Driver")
        self.assertEqual(report["drivers"][1]["gross_pay"], "200.00")

    def test_review_rejects_unassigned_trip_after_showing_route_earnings(self):
        client = FakeSupabaseClient()
        manager = FleetManager(client=client)
        client.tables["trips"] = [{
            "id": 12,
            "driver_id": None,
            "driver_name": "Manifest Only Driver",
            "ism_no": "ISM-IMPORT",
            "shipment_date": "2026-10-03",
            "origin": "GENSAN",
            "destination": "PLAS",
            "time_in": None,
            "time_out": None,
        }]

        report = manager.payroll_report("2026-10-01", "2026-10-15")
        self.assertEqual(report["drivers"][0]["gross_pay"], "900.00")
        with self.assertRaisesRegex(ValueError, "registered driver"):
            manager.review_payroll(
                "2026-10-01", "2026-10-15",
                "00000000-0000-0000-0000-000000000001",
            )

    def test_payroll_counts_configured_route_without_trip_times(self):
        client = FakeSupabaseClient()
        manager = FleetManager(client=client)
        vehicle = manager.add_vehicle("PAY-INCOMPLETE", "Toyota", "Hiace", 2022)
        driver = manager.add_driver("Incomplete Driver", "PAY-DL-INCOMPLETE")
        manager.add_trip(
            vehicle["id"], driver["id"], "GENSAN to PLAS", 0, 0,
            ism_no="ISM-INCOMPLETE", shipment_date="2026-10-03",
            time_in="08:00", origin="GENSAN", destination="PLAS",
        )

        report = manager.payroll_report("2026-10-01", "2026-10-15")

        self.assertEqual(report["trips"][0]["rate"], "900.00")
        self.assertTrue(report["trips"][0]["eligible"])
        self.assertIsNone(report["trips"][0]["issue"])
        self.assertEqual(report["drivers"][0]["gross_pay"], "900.00")

    def test_review_snapshots_eligible_trip_rates(self):
        client = FakeSupabaseClient()
        manager = FleetManager(client=client)
        vehicle = manager.add_vehicle("PAY-002", "Toyota", "Hiace", 2022)
        driver = manager.add_driver("Reviewed Driver", "PAY-DL-2")
        manager.add_trip(
            vehicle["id"], driver["id"], "GENSAN to DAV2", 0, 0,
            ism_no="ISM-900", shipment_date="2026-10-05",
            time_in="08:00", time_out="16:00",
            origin="GENSAN", destination="DAV2",
        )

        report = manager.review_payroll(
            "2026-10-01", "2026-10-15", "00000000-0000-0000-0000-000000000001"
        )

        self.assertEqual(report["status"], "reviewed")
        self.assertEqual(report["trips"][0]["rate"], "900.00")
        self.assertEqual(report["drivers"][0]["gross_pay"], "900.00")
        self.assertEqual(client.tables["payroll_items"][0]["trip_id"], 1)

    def test_payroll_report_api_returns_existing_trip_line_items(self):
        client = FakeSupabaseClient()
        manager = FleetManager(client=client)
        vehicle = manager.add_vehicle("PAY-003", "Toyota", "Hiace", 2022)
        driver = manager.add_driver("API Payroll Driver", "PAY-DL-3")
        manager.add_trip(
            vehicle["id"], driver["id"], "DAV1 to DAV2", 0, 0,
            ism_no="ISM-API", shipment_date="2026-10-06",
            time_in="08:00", time_out="11:00",
            origin="DAV1", destination="DAV2",
        )
        with patch.object(web_app, "manager", manager), patch.dict(
            web_app.app.config, {"TESTING": True, "TEST_AUTH_BYPASS": True}
        ):
            response = web_app.app.test_client().get(
                "/api/payroll?start_date=2026-10-01&end_date=2026-10-15"
            )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json["trips"][0]["ism_no"], "ISM-API")
        self.assertEqual(response.json["trips"][0]["rate"], "200.00")


class FakeAuthApiError(Exception):
    def __init__(self, message, status):
        super().__init__(message)
        self.status = status
        self.code = "invalid_credentials"


class FakeQuery:
    def __init__(self, client, table):
        self.client = client
        self.table_name = table
        self.action = "select"
        self.columns = "*"
        self.values = None
        self.filters = []
        self.sort = None
        self.row_limit = None
        self.row_range = None

    def select(self, columns="*", **kwargs):
        self.columns = columns
        return self

    def insert(self, values):
        self.action = "insert"
        self.values = values
        return self

    def update(self, values):
        self.action = "update"
        self.values = values
        return self

    def delete(self):
        self.action = "delete"
        return self

    def eq(self, column, value):
        self.filters.append((column, value))
        return self

    def order(self, column, desc=False):
        self.sort = (column, desc)
        return self

    def limit(self, count):
        self.row_limit = count
        return self

    def range(self, start, end):
        self.row_range = (start, end)
        return self

    def _matches(self, row):
        return all(row.get(key) == value for key, value in self.filters)

    def execute(self):
        rows = self.client.tables.setdefault(self.table_name, [])
        if self.action == "insert":
            inserted = []
            values = self.values if isinstance(self.values, list) else [self.values]
            for value in values:
                row = dict(value)
                row["id"] = self.client.next_ids.get(self.table_name, 0) + 1
                self.client.next_ids[self.table_name] = row["id"]
                if self.table_name == "payroll_periods":
                    row.setdefault("status", "draft")
                rows.append(row)
                inserted.append(dict(row))
            return SimpleNamespace(data=inserted)
        if self.action == "update":
            updated = []
            for row in rows:
                if self._matches(row):
                    row.update(self.values)
                    updated.append(row)
            return SimpleNamespace(data=[dict(row) for row in updated])
        if self.action == "delete":
            self.client.tables[self.table_name] = [
                row for row in rows if not self._matches(row)
            ]
            return SimpleNamespace(data=[])

        result = [dict(row) for row in rows if self._matches(row)]
        if self.sort:
            column, desc = self.sort
            result.sort(key=lambda row: row.get(column) or "", reverse=desc)
        if self.row_limit is not None:
            result = result[: self.row_limit]
        if self.row_range is not None:
            start, end = self.row_range
            result = result[start : end + 1]
        if (
            self.columns not in ("*",)
            and "driver:" not in self.columns
            and "vehicle:" not in self.columns
        ):
            selected = [column.strip() for column in self.columns.split(",")]
            result = [{key: row.get(key) for key in selected} for row in result]
        for row in result:
            if "driver:drivers(" in self.columns:
                driver = next(
                    (
                        item
                        for item in self.client.tables.get("drivers", [])
                        if item["id"] == row["driver_id"]
                    ),
                    None,
                )
                row["driver"] = {"name": driver["name"]} if driver else None
            if "vehicle:vehicles(" in self.columns:
                vehicle = next(
                    (
                        item
                        for item in self.client.tables.get("vehicles", [])
                        if item["id"] == row["vehicle_id"]
                    ),
                    None,
                )
                if vehicle:
                    row["vehicle"] = {
                        key: vehicle.get(key)
                        for key in ("plate_number", "make", "model")
                    }
                else:
                    row["vehicle"] = None
        return SimpleNamespace(data=result)


class FakeSupabaseClient:
    def __init__(self):
        self.tables = {}
        self.next_ids = {}

    def table(self, table_name):
        return FakeQuery(self, table_name)


class FakeStorage:
    def __init__(self):
        self.files = {}

    def from_(self, _bucket):
        return self

    def upload(self, path, file, _options):
        self.files[path] = file

    def download(self, path):
        return self.files[path]

    def remove(self, paths):
        for path in paths:
            self.files.pop(path, None)


class FakeAuthClient:
    def __init__(self):
        self.auth = self
        self.user = SimpleNamespace(id="auth-user-1", email="fleet@example.com")
        self.auth_session = SimpleNamespace(
            access_token="valid-access-token",
            refresh_token="valid-refresh-token",
            expires_in=3600,
        )

    def sign_in_with_password(self, credentials):
        if credentials != {"email": "fleet@example.com", "password": "correct-password"}:
            raise FakeAuthApiError("Invalid credentials", 400)
        return SimpleNamespace(session=self.auth_session, user=self.user)

    def get_user(self, access_token):
        if access_token != self.auth_session.access_token:
            raise ValueError("Expired access token")
        return SimpleNamespace(user=self.user)

    def refresh_session(self, refresh_token):
        if refresh_token != self.auth_session.refresh_token:
            raise ValueError("Invalid refresh token")
        return SimpleNamespace(session=self.auth_session, user=self.user)

class FleetManagerTests(unittest.TestCase):
    def setUp(self):
        self.client = FakeSupabaseClient()
        self.manager = FleetManager(client=self.client)

    def test_add_vehicle_driver_and_assignment(self):
        vehicle = self.manager.add_vehicle("ABC-123", "Toyota", "Hiace", 2022, "White")
        driver = self.manager.add_driver("Jane Doe", "DL-1001", phone="555-0100")
        assignment = self.manager.assign_vehicle(
            driver["id"], vehicle["id"], "2026-10-03"
        )
        listed = self.manager.list_assignments()

        self.assertEqual(vehicle["plate_number"], "ABC-123")
        self.assertEqual(driver["name"], "Jane Doe")
        self.assertEqual(assignment["driver_id"], driver["id"])
        self.assertEqual(listed[0]["driver_name"], "Jane Doe")
        self.assertEqual(listed[0]["plate_number"], "ABC-123")
        self.assertEqual(self.manager.get_vehicle(vehicle["id"])["status"], "assigned")

    def test_trip_and_maintenance_and_fuel_tracking(self):
        vehicle = self.manager.add_vehicle("XYZ-988", "Mercedes", "Sprinter", 2021, "Blue")
        driver = self.manager.add_driver("John Smith", "DL-9988")

        trip = self.manager.add_trip(
            vehicle["id"], driver["id"], "Accra to Kumasi", 15000, 15375,
            "2026-10-02", "Routine delivery route",
        )
        maintenance = self.manager.add_maintenance(
            vehicle["id"], "service", "Oil change", 120.0
        )
        fuel = self.manager.add_fuel(vehicle["id"], "diesel", 80, 15.5)

        self.assertEqual(trip["route"], "Accra to Kumasi")
        self.assertEqual(maintenance["service_type"], "service")
        self.assertEqual(fuel["total_cost"], 1240.0)
        self.assertEqual(self.manager.dashboard()["recent_trips"], 1)
        self.assertEqual(self.manager.dashboard()["total_fuel_spend"], 1240.0)
        self.assertEqual(self.manager.get_vehicle(vehicle["id"])["odometer"], 15375)

    def test_imported_trip_without_vehicle_or_driver_ids_displays_source_fields(self):
        self.client.tables["trips"] = [{
            "id": 1,
            "vehicle_id": None,
            "driver_id": None,
            "record_date": "2026-09-28",
            "ism_no": "31946326",
            "shipment_date": "2026-09-27",
            "time_in": "10:27:00",
            "time_out": None,
            "origin": "GENSAN",
            "destination": "DAV2",
            "plate_no": "MUB289",
            "load_details": "1P",
            "driver_name": "SALCEDA",
            "trip_fuel": "1ST-105 INV12125 9.27",
            "route": "GENSAN to DAV2",
            "trip_date": "2026-09-27",
        }]

        trip = self.manager.list_trips()[0]
        self.assertEqual(trip["plate_number"], "MUB289")
        self.assertEqual(trip["driver_name"], "SALCEDA")

    def test_export_report(self):
        self.manager.add_vehicle("ZZZ-001", "Ford", "Transit", 2023)
        self.manager.add_driver("Esther Kofi", "DL-4411")
        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = os.path.join(temp_dir, "report.xlsx")
            result = self.manager.export_report(output_path)
            self.assertTrue(os.path.exists(result["output_path"]))
            self.assertGreater(result["records_written"], 0)
            workbook = load_workbook(output_path, data_only=True)
            self.assertEqual(workbook["Vehicles"]["B5"].value, "ZZZ-001")

    def test_web_export_returns_preformatted_system_wide_excel_workbook(self):
        vehicle = self.manager.add_vehicle("XLS-001", "Toyota", "Hiace", 2024)
        driver = self.manager.add_driver("Excel Driver", "XLS-DL-1")
        self.manager.assign_vehicle(driver["id"], vehicle["id"])
        self.manager.add_maintenance(
            vehicle["id"], "Service", '=HYPERLINK("https://example.com","Open")', 1250.5
        )
        self.manager.add_fuel(vehicle["id"], "Diesel", 50, 60.25)
        self.manager.add_trip(
            vehicle["id"], driver["id"], "Manila to Cavite", 100, 150,
            origin="Manila", destination="Cavite",
        )
        web_app.manager = self.manager
        web_app.app.config["TESTING"] = True
        client = web_app.app.test_client()
        self.sign_in(client)

        response = client.get("/api/export")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.mimetype,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn("fleet_system_report.xlsx", response.headers["Content-Disposition"])
        workbook = load_workbook(BytesIO(response.data), data_only=True)
        self.assertEqual(
            workbook.sheetnames,
            ["Dashboard", "Vehicles", "Drivers", "Assignments", "Maintenance", "Fuel Logs", "Trips"],
        )
        vehicles = workbook["Vehicles"]
        self.assertEqual(vehicles["B4"].value, "Plate Number")
        self.assertEqual(vehicles["B5"].value, "XLS-001")
        self.assertEqual(vehicles.freeze_panes, "A5")
        self.assertEqual(vehicles.auto_filter.ref, "A4:K5")
        self.assertEqual(workbook["Fuel Logs"]["G5"].value, 3012.5)
        self.assertIn("₱", workbook["Fuel Logs"]["G5"].number_format)
        self.assertEqual(workbook["Maintenance"]["E5"].data_type, "s")
        self.assertEqual(workbook["Trips"]["G5"].value, "Manila")

    def test_web_trip_history_export_returns_trip_details_as_excel(self):
        vehicle = self.manager.add_vehicle("TRIP-XLS-1", "Toyota", "Hiace", 2024)
        driver = self.manager.add_driver("Trip Export Driver", "TRIP-DL-1")
        self.manager.add_trip(
            vehicle["id"],
            driver["id"],
            "Manila to Cavite",
            100,
            150,
            ism_no="ISM-TRIP-1",
            shipment_date="2026-10-05",
            time_in="08:15",
            time_out="10:30",
            origin="Manila",
            destination="Cavite",
            load_details="12 cartons",
            trip_fuel="Diesel",
        )
        web_app.manager = self.manager
        web_app.app.config["TESTING"] = True
        client = web_app.app.test_client()
        self.sign_in(client)

        response = client.get("/api/export/trips")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.mimetype,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn("trip_history.xlsx", response.headers["Content-Disposition"])
        workbook = load_workbook(BytesIO(response.data), data_only=True)
        self.assertEqual(workbook.sheetnames, ["Trip History"])
        trips = workbook["Trip History"]
        self.assertEqual(trips["B4"].value, "Ism No")
        self.assertEqual(trips["B5"].value, "ISM-TRIP-1")
        self.assertEqual(trips["F5"].value, "Manila")
        self.assertEqual(trips["H5"].value, "TRIP-XLS-1")
        self.assertEqual(trips["K5"].value, "Trip Export Driver")
        self.assertEqual(trips["C5"].number_format, "yyyy-mm-dd")
        self.assertEqual(trips["D5"].number_format, "hh:mm")
        self.assertEqual(trips.freeze_panes, "A5")
        self.assertEqual(trips.auto_filter.ref, "A4:R5")

    def test_web_trip_history_export_uses_selected_group_and_search(self):
        vehicle = self.manager.add_vehicle("TRIP-XLS-2", "Toyota", "Hiace", 2024)
        first_driver = self.manager.add_driver("Alpha Driver", "TRIP-DL-2")
        second_driver = self.manager.add_driver("Beta Driver", "TRIP-DL-3")
        self.manager.add_trip(
            vehicle["id"],
            first_driver["id"],
            "Manila to Cavite",
            100,
            150,
            ism_no="MATCH-TRIP",
            origin="Manila",
            destination="Cavite",
        )
        self.manager.add_trip(
            vehicle["id"],
            second_driver["id"],
            "Cavite to Manila",
            150,
            200,
            ism_no="OTHER-TRIP",
            origin="Cavite",
            destination="Manila",
        )
        web_app.manager = self.manager
        web_app.app.config["TESTING"] = True
        client = web_app.app.test_client()
        self.sign_in(client)

        response = client.get(
            "/api/export/trips?group_by=driver_name&search=MATCH-TRIP"
        )

        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.data), data_only=True)
        self.assertEqual(workbook.sheetnames, ["Alpha Driver"])
        self.assertEqual(workbook["Alpha Driver"]["B5"].value, "MATCH-TRIP")

    def test_cli_export_report_defaults_to_excel_file(self):
        from fleet_management import _build_cli

        args = _build_cli().parse_args(["export-report"])
        self.assertEqual(args.output, "fleet_system_report.xlsx")

    def test_missing_supabase_configuration_is_reported(self):
        with patch.dict(os.environ, {}, clear=True):
            with self.assertRaisesRegex(FleetManagementError, "SUPABASE_URL"):
                FleetManager()

    def test_supabase_rest_endpoint_is_normalized_to_project_url(self):
        self.assertEqual(
            _supabase_project_url("https://project-ref.supabase.co/rest/v1/"),
            "https://project-ref.supabase.co",
        )

    def test_supabase_client_receives_normalized_url(self):
        with patch("fleet_management.create_client") as create_client:
            FleetManager(
                url="https://project-ref.supabase.co/rest/v1/",
                key="test-key",
            )
        create_client.assert_called_once_with(
            "https://project-ref.supabase.co", "test-key"
        )

    def test_dns_error_nested_in_transport_exception_gets_actionable_message(self):
        class BrokenQuery:
            def execute(self):
                try:
                    raise OSError(11001, "getaddrinfo failed")
                except OSError as cause:
                    raise RuntimeError("request transport failed") from cause

        with patch.dict(
            os.environ,
            {"SUPABASE_URL": "https://project-ref.supabase.co"},
        ):
            with self.assertRaisesRegex(
                FleetManagementError, "Could not resolve the Supabase server"
            ):
                self.manager._execute(BrokenQuery())

    def test_missing_trip_migration_has_actionable_error(self):
        class UnmigratedTripsQuery:
            def execute(self):
                raise RuntimeError(
                    "{'message': 'column trips.shipment_date does not exist', "
                    "'code': '42703'}"
                )

        with self.assertRaisesRegex(
            FleetManagementError, "supabase_trips_migration.sql"
        ):
            self.manager._execute(UnmigratedTripsQuery())

    def test_web_dashboard_and_vehicle_form_use_supabase_manager(self):
        web_app.manager = self.manager
        web_app.app.config["TESTING"] = True
        client = web_app.app.test_client()

        self.sign_in(client)
        response = client.get("/")
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"<!doctype html>", response.data.lower())
        self.assertIn(b"/static/app.js", response.data)
        response.close()

        response = client.post(
            "/api/vehicles",
            json={
                "plate_number": "WEB-123",
                "make": "Toyota",
                "model": "Hiace",
                "year": 2022,
            },
            headers={"X-CSRF-Token": self.csrf_token(client)},
        )
        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.json["data"]["plate_number"], "WEB-123")

        driver = self.manager.add_driver("Web Driver", "WEB-DL-1")
        vehicle = self.manager.list_vehicles()[0]
        response = client.post(
            "/api/trips",
            json={
                "ism_no": "ISM-008",
                "shipment_date": "2026-10-03",
                "time_in": "08:15",
                "time_out": "09:40",
                "origin": "Manila",
                "destination": "Cavite",
                "vehicle_id": vehicle["id"],
                "driver_id": driver["id"],
                "load_details": "12 cartons",
                "trip_fuel": "Diesel",
            },
            headers={"X-CSRF-Token": self.csrf_token(client)},
        )
        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.json["data"]["ism_no"], "ISM-008")
        self.assertEqual(response.json["data"]["origin"], "Manila")
        self.assertEqual(response.json["data"]["time_in"], "08:15")
        self.assertEqual(response.json["data"]["trip_fuel"], "Diesel")

        response = client.get("/api/bootstrap")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json["vehicles"][0]["plate_number"], "WEB-123")
        trip = response.json["trips"][0]
        self.assertEqual(trip["shipment_date"], "2026-10-03")
        self.assertEqual(trip["plate_number"], "WEB-123")
        self.assertEqual(trip["driver_name"], "Web Driver")

    def test_dispatcher_is_limited_to_owned_records_and_allowed_endpoints(self):
        web_app.manager = self.manager
        web_app.app.config["TESTING"] = True
        vehicle = self.manager.add_vehicle("DSP-123", "Toyota", "Hiace", 2022)
        driver = self.manager.add_driver("Dispatcher Driver", "DSP-DL-1")
        own_trip = self.manager.add_trip(
            vehicle["id"], driver["id"], "Origin to Destination", 0, 0,
            ism_no="DSP-OWN", origin="Origin", destination="Destination",
            created_by="auth-user-1",
        )
        self.manager.add_trip(
            vehicle["id"], driver["id"], "Other to Route", 0, 0,
            ism_no="OTHER-TRIP", origin="Other", destination="Route",
            created_by="another-dispatcher",
        )
        auth_client = FakeAuthClient()
        auth_client.user.user_metadata = {"role": "dispatcher", "full_name": "Dispatch User"}
        auth_client.user.app_metadata = {"role": "administrator"}
        client = web_app.app.test_client()
        self.sign_in(client, auth_client)

        bootstrap = client.get("/api/bootstrap")
        self.assertEqual(bootstrap.status_code, 200)
        self.assertEqual(bootstrap.json["user"]["role"], "dispatcher")
        self.assertEqual([row["id"] for row in bootstrap.json["trips"]], [own_trip["id"]])
        self.assertEqual(bootstrap.json["dashboard"]["total_trips"], 1)
        self.assertNotIn("total_vehicles", bootstrap.json["dashboard"])
        self.assertEqual(
            bootstrap.json["drivers"],
            [{"id": driver["id"], "name": "Dispatcher Driver", "status": "active"}],
        )

        csrf = bootstrap.json["csrf_token"]
        denied_vehicle_create = client.post(
            "/api/vehicles",
            json={"plate_number": "DENIED", "make": "Toyota", "model": "Hiace", "year": 2022},
            headers={"X-CSRF-Token": csrf},
        )
        denied_vehicle_list = client.get("/api/vehicles")
        denied_driver_list = client.get("/api/drivers")
        denied_payroll = client.get("/api/payroll?start_date=2026-10-01&end_date=2026-10-31")
        denied_export = client.get("/api/export")
        denied_delete = client.delete(
            f"/api/trips/{own_trip['id']}",
            headers={"X-CSRF-Token": csrf},
        )
        self.assertEqual(denied_vehicle_create.status_code, 403)
        self.assertEqual(denied_vehicle_list.status_code, 403)
        self.assertEqual(denied_driver_list.status_code, 403)
        self.assertEqual(denied_payroll.status_code, 403)
        self.assertEqual(denied_export.status_code, 403)
        self.assertEqual(denied_delete.status_code, 403)

        created_trip = client.post(
            "/api/trips",
            json={
                "ism_no": "DSP-NEW",
                "shipment_date": "2026-10-03",
                "origin": "Origin",
                "destination": "Destination",
                "vehicle_id": vehicle["id"],
                "driver_id": driver["id"],
                "created_by": "another-dispatcher",
            },
            headers={"X-CSRF-Token": csrf},
        )
        self.assertEqual(created_trip.status_code, 400)
        self.assertEqual(len(self.manager.list_trips(user_id="auth-user-1")), 1)

    def test_dispatcher_can_upload_and_view_owned_trip_manifest_and_fuel_invoice(self):
        web_app.manager = self.manager
        web_app.app.config["TESTING"] = True
        self.manager.client.storage = FakeStorage()
        vehicle = self.manager.add_vehicle("DOC-123", "Toyota", "Hiace", 2022)
        driver = self.manager.add_driver("Document Driver", "DOC-DL-1")
        trip = self.manager.add_trip(
            vehicle["id"], driver["id"], "Origin to Destination", 0, 0,
            ism_no="DOC-TRIP", origin="Origin", destination="Destination",
            created_by="auth-user-1",
        )
        other_trip = self.manager.add_trip(
            vehicle["id"], driver["id"], "Other to Destination", 0, 0,
            ism_no="OTHER-DOC-TRIP", origin="Other", destination="Destination",
            created_by="another-dispatcher",
        )
        fuel = self.manager.add_fuel(
            vehicle["id"], "diesel", 20, 1.5, created_by="auth-user-1"
        )

        auth_client = FakeAuthClient()
        auth_client.user.user_metadata = {"role": "dispatcher"}
        client = web_app.app.test_client()
        self.sign_in(client, auth_client)
        csrf = self.csrf_token(client)

        image = BytesIO()
        Image.new("RGB", (2, 2), color="blue").save(image, format="PNG")
        image_bytes = image.getvalue()
        manifest_upload = client.post(
            f"/api/trips/{trip['id']}/manifest",
            data={"image": (BytesIO(image_bytes), "manifest.png")},
            headers={"X-CSRF-Token": csrf},
        )
        invoice_upload = client.post(
            f"/api/fuel/{fuel['id']}/invoice",
            data={"image": (BytesIO(image_bytes), "invoice.png")},
            headers={"X-CSRF-Token": csrf},
        )
        bootstrap = client.get("/api/bootstrap")
        manifest_view = client.get(f"/api/trips/{trip['id']}/manifest")
        invoice_view = client.get(f"/api/fuel/{fuel['id']}/invoice")
        other_manifest_view = client.get(f"/api/trips/{other_trip['id']}/manifest")

        self.assertEqual(manifest_upload.status_code, 201)
        self.assertEqual(invoice_upload.status_code, 201)
        self.assertEqual(
            bootstrap.json["trips"][0]["manifest_image_url"],
            f"/api/trips/{trip['id']}/manifest",
        )
        self.assertTrue(bootstrap.json["fuel_logs"][0]["invoice_image_url"].endswith("/invoice"))
        self.assertEqual(manifest_view.status_code, 200)
        self.assertEqual(invoice_view.status_code, 200)
        self.assertEqual(manifest_view.mimetype, "image/png")
        self.assertEqual(manifest_view.data, image_bytes)
        self.assertEqual(invoice_view.data, image_bytes)
        self.assertEqual(other_manifest_view.status_code, 403)

        other_manifest_upload = client.post(
            f"/api/trips/{other_trip['id']}/manifest",
            data={"image": (BytesIO(image_bytes), "manifest.png")},
            headers={"X-CSRF-Token": csrf},
        )
        self.assertEqual(other_manifest_upload.status_code, 403)

        invalid_image = client.post(
            f"/api/trips/{trip['id']}/manifest",
            data={"image": (BytesIO(b"not an image"), "manifest.png")},
            headers={"X-CSRF-Token": csrf},
        )
        self.assertEqual(invalid_image.status_code, 400)

    def test_web_reports_missing_supabase_configuration(self):
        web_app.manager = None
        web_app.app.config["TESTING"] = True
        client = web_app.app.test_client()
        self.sign_in(client)
        with patch.dict(os.environ, {}, clear=True):
            response = client.get("/api/bootstrap")
        self.assertEqual(response.status_code, 503)
        self.assertIn(b"SUPABASE_URL", response.data)

    def csrf_token(self, client):
        response = client.get("/api/bootstrap")
        self.assertEqual(response.status_code, 200)
        return response.json["csrf_token"]

    def sign_in(self, client, auth_client=None):
        patcher = patch.object(
            web_app,
            "get_auth_client",
            return_value=auth_client or FakeAuthClient(),
        )
        patcher.start()
        self.addCleanup(patcher.stop)
        csrf = client.get("/api/auth/csrf").json["csrf_token"]
        response = client.post(
            "/api/auth/login",
            json={"email": "fleet@example.com", "password": "correct-password"},
            headers={"X-CSRF-Token": csrf},
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("HttpOnly", response.headers.get("Set-Cookie", ""))
        cookies = response.headers.getlist("Set-Cookie")
        for name in ("glfs_access_token", "glfs_refresh_token"):
            auth_cookie = next(cookie for cookie in cookies if cookie.startswith(f"{name}="))
            self.assertNotIn("Max-Age=", auth_cookie)
            self.assertNotIn("Expires=", auth_cookie)
        with client.session_transaction() as browser_session:
            self.assertFalse(browser_session.permanent)
            self.assertIn("last_activity", browser_session)
        return response

    def test_app_and_api_require_authentication_but_login_assets_are_public(self):
        web_app.app.config["TESTING"] = True
        client = web_app.app.test_client()

        self.assertEqual(client.get("/").status_code, 302)
        self.assertEqual(client.get("/api/bootstrap").status_code, 401)
        self.assertEqual(client.get("/static/index.html").status_code, 302)
        login_page = client.get("/login")
        self.assertEqual(login_page.status_code, 200)
        login_page.close()
        login_script = client.get("/static/login.js")
        self.assertEqual(login_script.status_code, 200)
        login_script.close()

    def test_login_rejects_bad_credentials_and_requires_csrf(self):
        web_app.app.config["TESTING"] = True
        client = web_app.app.test_client()
        with patch.object(web_app, "get_auth_client", return_value=FakeAuthClient()):
            csrf = client.get("/api/auth/csrf").json["csrf_token"]
            no_csrf = client.post(
                "/api/auth/login",
                json={"email": "fleet@example.com", "password": "correct-password"},
            )
            bad_password = client.post(
                "/api/auth/login",
                json={"email": "fleet@example.com", "password": "wrong-password"},
                headers={"X-CSRF-Token": csrf},
            )
        self.assertEqual(no_csrf.status_code, 403)
        self.assertEqual(bad_password.status_code, 401)
        self.assertIn(b"Sign-in failed", bad_password.data)

    def test_login_reports_missing_supabase_anon_key(self):
        web_app.app.config["TESTING"] = True
        client = web_app.app.test_client()
        with patch.object(web_app, "_supabase_setting", return_value=None):
            csrf = client.get("/api/auth/csrf").json["csrf_token"]
            response = client.post(
                "/api/auth/login",
                json={"email": "fleet@example.com", "password": "correct-password"},
                headers={"X-CSRF-Token": csrf},
            )
        self.assertEqual(response.status_code, 503)
        self.assertIn(b"SUPABASE_ANON_KEY", response.data)

    def test_mutations_require_csrf_and_logout_clears_authenticated_session(self):
        web_app.manager = self.manager
        web_app.app.config["TESTING"] = True
        client = web_app.app.test_client()
        auth_client = FakeAuthClient()
        self.sign_in(client, auth_client)
        csrf = self.csrf_token(client)
        payload = {
            "plate_number": "CSRF-1",
            "make": "Toyota",
            "model": "Hiace",
            "year": 2022,
        }
        blocked = client.post("/api/vehicles", json=payload)
        created = client.post(
            "/api/vehicles",
            json=payload,
            headers={"X-CSRF-Token": csrf},
        )
        logout = client.post(
            "/api/auth/logout",
            headers={"X-CSRF-Token": csrf},
        )
        protected = client.get("/api/bootstrap")
        self.assertEqual(blocked.status_code, 403)
        self.assertEqual(created.status_code, 201)
        self.assertEqual(logout.status_code, 200)
        self.assertEqual(protected.status_code, 401)

    def test_logout_clears_local_session_without_csrf_token(self):
        web_app.app.config["TESTING"] = True
        client = web_app.app.test_client()
        self.sign_in(client)

        response = client.post(
            "/api/auth/logout",
        )

        self.assertEqual(response.status_code, 200)
        self.assertIn("/login", response.json["redirect"])
        self.assertTrue(any(
            "glfs_access_token=;" in cookie
            for cookie in response.headers.getlist("Set-Cookie")
        ))
        self.assertEqual(client.get("/api/bootstrap").status_code, 401)

    def test_expired_access_token_is_refreshed(self):
        web_app.app.config["TESTING"] = True
        client = web_app.app.test_client()
        auth_client = FakeAuthClient()
        self.sign_in(client, auth_client)
        auth_client.auth_session.access_token = "refreshed-access-token"
        # The old browser token expires; the refresh token supplies a new one.
        response = client.get("/api/auth/me")
        self.assertEqual(response.status_code, 200)
        self.assertIn("refreshed-access-token", response.headers.get("Set-Cookie", ""))

    def test_idle_session_expires_after_three_minutes_and_clears_auth_cookies(self):
        web_app.app.config["TESTING"] = True
        client = web_app.app.test_client()
        self.sign_in(client)
        with client.session_transaction() as browser_session:
            browser_session["last_activity"] = (
                time.time() - web_app.SESSION_IDLE_TIMEOUT_SECONDS - 1
            )

        response = client.get("/api/bootstrap")

        self.assertEqual(response.status_code, 401)
        self.assertIn(b"inactivity", response.data)
        self.assertTrue(any(
            "glfs_refresh_token=;" in cookie
            for cookie in response.headers.getlist("Set-Cookie")
        ))

    def test_idle_browser_redirect_shows_inactivity_notice(self):
        web_app.app.config["TESTING"] = True
        client = web_app.app.test_client()
        self.sign_in(client)
        with client.session_transaction() as browser_session:
            browser_session["last_activity"] = (
                time.time() - web_app.SESSION_IDLE_TIMEOUT_SECONDS - 1
            )

        response = client.get("/")

        self.assertEqual(response.status_code, 302)
        self.assertIn("/login?expired=1", response.headers["Location"])
        login_page = client.get(response.headers["Location"])
        self.assertIn(b"You have been logged out due to inactivity.", login_page.data)
        login_page.close()

    def test_expired_browser_redirect_forces_login_page(self):
        web_app.app.config["TESTING"] = True
        client = web_app.app.test_client()
        self.sign_in(client)

        response = client.get("/login?expired=1")

        self.assertEqual(response.status_code, 200)
        self.assertIn(b"Welcome back", response.data)
        self.assertTrue(any(
            "glfs_access_token=;" in cookie
            for cookie in response.headers.getlist("Set-Cookie")
        ))
        response.close()


if __name__ == "__main__":
    unittest.main()
